"""Read supplied sources without modifying them; rebuild site content/data assets."""
from pathlib import Path
import json, shutil, re, html
from zipfile import ZipFile
from lxml import etree as E
import openpyxl
root=Path(__file__).resolve().parents[2]; site=root/'mantle-modeling-portfolio'
w=openpyxl.load_workbook(root/'figures/melting_xtal_plot_data.xlsx',data_only=True,read_only=True)
data={}
for kind,sheet,names,start in [('melting','melting_plot',['Olivine','Orthopyroxene','Clinopyroxene','Garnet','Spinel'],6),('crystallization','crystallization_plot',['Olivine','Plagioclase','Clinopyroxene'],4)]:
 rows=[r for r in w[sheet].iter_rows(min_row=2,values_only=True) if isinstance(r[0],(int,float))]
 col=lambda i:[r[i] for r in rows]
 degree=col(0)
 d={'degree':[v/100 for v in degree],'melt':[v/100 for v in col(0 if kind=='melting' else 17)],'minerals':[{'name':n,'values':[v/100 for v in col(i+1)]} for i,n in enumerate(names)],'proportions':[{'name':n,'values':col(start+i)} for i,n in enumerate(names)],'charts':[]}
 for i,(el,unit) in enumerate([('MgO','wt%'),('La','ppm')]):
  label=f'{"Melt" if kind=="melting" else "Residual magma"} {el} ({unit})'
  d['charts'].append({'title':label,'yLabel':label,'series':[{'name':el,'values':col(start+len(names)+i)}]})
 if kind=='melting':
  d['initialState']=True
  d['degree'].insert(0,0);d['melt'].insert(0,0)
  for m,v in zip(d['minerals'],[.525,.15,.25,.075,0]):m['values'].insert(0,v)
  for m,v in zip(d['proportions'],[52.5,15,25,7.5,0]):m['values'].insert(0,v)
  for chart in d['charts']:
   for series in chart['series']:series['values'].insert(0,None)
 if kind=='crystallization':
  d['initialState']=True
  d['degree'].insert(0,0);d['melt'].insert(0,1)
  for m in d['minerals']:m['values'].insert(0,0)
  for m in d['proportions']:m['values'].insert(0,None)
  for chart in d['charts']:
   for series in chart['series']:series['values'].insert(0,None)
 data[kind]=d
 print(kind,len(rows),'max phase total error',max(abs((sum(r[:6]) if kind=='melting' else r[17]+sum(r[1:4]))-100) for r in rows))
(site/'model-data.js').write_text('/* Extracted from figures/melting_xtal_plot_data.xlsx. See scripts/import-content.py. */\nwindow.MANTLE_MODEL_DATA = '+json.dumps(data,separators=(',',':'))+';\n')
# Convert Word's structured equations to native MathML, retaining fractions, subscripts and sums.
r=E.fromstring(ZipFile(root/'website idea.docx').read('word/document.xml'));ns={'w':r.nsmap['w'],'m':r.nsmap['m']}; paragraphs=r.findall('w:body/w:p',ns)
def math(e):
 tag=E.QName(e).localname
 children=lambda:''.join(math(c) for c in e)
 part=lambda name:'<mrow>'+''.join(math(c) for c in e.find('m:'+name,ns))+'</mrow>'
 if tag.endswith('Pr'):return ''
 if tag=='t':return ('<mspace width="2em"/>' if 'Eq.' in (e.text or '') else '')+'<mtext>'+html.escape(e.text or '')+'</mtext>'
 if tag=='f' and ''.join(e.find('m:num',ns).itertext())=='ol' and ''.join(e.find('m:den',ns).itertext())=='l':return '<mtext>ol/l</mtext>'
 if tag=='f':return '<mfrac>'+part('num')+part('den')+'</mfrac>'
 if tag in ('sSub','sSup','sSubSup'):
  t={'sSub':'msub','sSup':'msup','sSubSup':'msubsup'}[tag]
  return '<'+t+'>'+part('e')+('' if tag=='sSup' else part('sub'))+('' if tag=='sSub' else part('sup'))+'</'+t+'>'
 if tag=='nary':
  char=e.find('m:naryPr/m:chr',ns); symbol=char.get('{'+ns['m']+'}val') if char is not None else '∑'
  return '<munderover><mo>'+symbol+'</mo>'+part('sub')+part('sup')+'</munderover>'+part('e')
 if tag=='d':return '<mo>(</mo>'+part('e')+'<mo>)</mo>'
 return children()
def inline(e):
 tag=E.QName(e).localname
 if tag in ('drawing','pict'):return ''
 if tag=='oMath':return '<math xmlns="http://www.w3.org/1998/Math/MathML">'+math(e)+'</math>'
 if tag=='t':return html.escape(e.text or '')
 return ''.join(inline(c) for c in e)
def equation_markup(content):
 # Equation labels live outside the math in a shared grid column.
 label=re.search(r'Eq\.\s*([mc]\d[a-z]?\d?)',re.sub('<[^>]+>','',content))
 if not label:return content
 number=label[1]
 content=re.sub(r'<mspace[^>]*/>','',content)
 content=re.sub(r'<mtext>([^<]*?)Eq\.[^<]*</mtext>',lambda m:'<mtext>'+m[1].rstrip()+'</mtext>',content)
 content=re.sub(r'\s*Eq\.\s*[mc]\d[a-z]?\d?','',content)
 suffix=''
 tail=re.search(r'for every mineral ([ab])',re.sub('<[^>]+>','',content))
 if tail:
  suffix=' for every mineral '+tail[1]
  content=content[:content.index('</math>')+7]
 content=re.sub(r'<mtext>\s*</mtext>','',content)
 return '<div class="equation-row"><div class="equation-expression">'+content+'</div><span class="equation-number">Eq. '+number+suffix+'</span></div>'
def para(i):
 content=inline(paragraphs[i])
 if i in [7,8,9,12,14,16,17,22,23,48,49]:return '<div class="equations">'+equation_markup(content)+'</div>'
 if i==46:
  content=content.replace('is the extent of crystallization','is the remaining melt fraction')
  # The final math in this paragraph is Eq. c1; give its label the same spacing.
  maths=list(re.finditer(r'<math.*?</math>',content))
  last=maths[-1]
  return '<p>'+content[:last.start()]+' Crystallization degree is 1−F.</p><div class="equations">'+equation_markup(last[0])+'</div>'
 return '<p>'+content+'</p>'
h=(site/'index.html').read_text()
for kind,ids in [('melting',range(6,25)),('crystallization',range(46,51))]:
 content='\n'.join(para(i) for i in ids if inline(paragraphs[i]).strip())
 pattern=r'(<summary>Explore the '+kind+r' method.*?</summary>)<div class="method-content">.*?</div></details>'
 h=re.sub(pattern,lambda m:m[1]+'<div class="method-content">'+content+'</div></details>',h,flags=re.S)
h=h.replace('The initial source contains no spinel. A garnet-to-spinel conversion reaction changes the mineral proportions before the first calculated melting step.','The spinel phase was created from a garnet-to-spinel conversion reaction.')
h=h.replace('Source composition: Synthetic','Source composition: synthetic')
h=re.sub(r'(<figure class="animation-panel"[^>]*>)<figcaption>.*?</figcaption>',r'\1',h)
ni='<div class="ni-definition"><div class="equations">'+inline(paragraphs[58])+'</div><div class="ni-equation"><div class="equations">'+inline(paragraphs[59])+'</div>'+para(60)+'</div></div>'
h=re.sub(r'<div class="ni-definition">.*?</div></div></div>',ni+'</div>',h,flags=re.S)
h=h.replace('shows a better performance in having smaller relative root mean square error (RRMSE), higher R2, and smaller intercept, and slope closer one,','shows better performance, with a smaller relative root mean square error (RRMSE), higher R<sup>2</sup>, smaller intercept, and a slope closer to one,')
# Move the figure immediately after the experimental-comparison paragraph.
fig=re.search(r'<figure class="research-figure"><img src="assets/nickel-comparison.png".*?</figure>',h,re.S)
if fig:
 figure=fig[0];h=h.replace(figure,'')
 target=re.search(r'<p class="nickel-motivation">(?:When comparing|Compared with experimental).*?</p>',h,re.S)
 note='<p class="nickel-dataset-note">The details of the experimental dataset and comparisons are listed in <a href="https://doi.org/10.1016/j.chemgeo.2023.121745">Yu &amp; Langmuir (2023)</a>.</p>'
 h=re.sub(r'<p class="nickel-dataset-note">.*?</p>','',h,flags=re.S)
 h=h.replace(target[0],target[0]+figure+note)
h=h.replace('assets/melting-temperature.svg','assets/melting-mgo.svg').replace('alt="Modeled and experimental temperature"','alt="Modeled and experimental MgO"')
# Identify the second comparison independently to avoid filename substitution collisions.
h=re.sub(r'(<div class="comparison-figures"><h3>Comparison with experimental data.*?<figure.*?</figure><figure class="research-figure"><img src=")[^"]+(" alt=")[^"]+',r'\1assets/melting-k2o.svg\2Modeled and experimental K₂O',h,flags=re.S)
(site/'index.html').write_text(h)
# Independent comparison pairs retain their own row counts.
# Standalone SVG comparison plots, with independent valid numeric pairs per series.
def pairs(sheet,a,b):
 return [(r[a],r[b]) for r in w[sheet].iter_rows(min_row=2,values_only=True) if isinstance(r[a],(int,float)) and isinstance(r[b],(int,float))]
for file,sheet,a,b,c,d,xlabel,ylabel in [('melting-mgo','melting_plot',13,14,None,None,'Modeled MgO (wt%)','Experiment MgO (wt%)'),('melting-k2o','melting_plot',15,16,None,None,'Modeled K₂O (wt%)','Experiment K₂O (wt%)'),('crystallization-melt','crystallization_plot',9,10,11,12,'magma MgO (wt%)','magma total FeO (wt%)'),('crystallization-olivine','crystallization_plot',13,14,15,16,'olivine Fo','olivine Ni (ppm)')]:
 p=pairs(sheet,a,b);q=pairs(sheet,c,d) if c is not None else [];allp=p+q
 xs=[v[0] for v in allp];ys=[v[1] for v in allp];xmin,xmax=min(xs),max(xs);ymin,ymax=min(ys),max(ys)
 if c is None:xmin=ymin=min(xmin,ymin);xmax=ymax=max(xmax,ymax)
 dx=(xmax-xmin)*.07;dy=(ymax-ymin)*.07;xmin-=dx;xmax+=dx;ymin-=dy;ymax+=dy
 if file=='melting-mgo':xmin,xmax,ymin,ymax=15,25,15,25
 if file=='melting-k2o':xmin,xmax,ymin,ymax=.1,.7,.1,.7
 if file=='crystallization-melt':xmin,xmax,ymin,ymax=4.5,11,6.5,15
 if file=='crystallization-olivine':xmin,xmax,ymin,ymax=80,92,340,3245
 x=lambda v:76+(v-xmin)/(xmax-xmin)*480;y=lambda v:348-(v-ymin)/(ymax-ymin)*280
 svg=['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 600 440" role="img"><title>'+html.escape(ylabel+' versus '+xlabel)+'</title><rect width="600" height="440" fill="white"/><g font-family="Avenir Next, sans-serif" font-size="17" fill="#173033">']
 for i in range(5):
  vx=xmin+(xmax-xmin)*i/4;vy=ymin+(ymax-ymin)*i/4
  svg.append(f'<path d="M76 {y(vy)}H556" stroke="#e3e9e7"/><text x="66" y="{y(vy)+5}" text-anchor="end">{round(vy) if abs(vy)>=1000 else round(vy,2):g}</text><text x="{x(vx)}" y="372" text-anchor="middle">{round(vx) if abs(vx)>=1000 else round(vx,2):g}</text>')
 svg.append('<path d="M76 68V348H556" fill="none" stroke="#82958f"/>')
 if c is None:svg.append(f'<path d="M{x(xmin)} {y(xmin)}L{x(xmax)} {y(xmax)}" stroke="#8b9894" stroke-dasharray="5 5"/>')
 svg.append('<defs><clipPath id="plot"><rect x="76" y="68" width="480" height="280"/></clipPath></defs><g clip-path="url(#plot)">')
 for vx,vy in p:svg.append(f'<circle cx="{x(vx):.3f}" cy="{y(vy):.3f}" r="{6 if c is None else 2.8}" fill="{"#46658f" if c is None else "#657da9"}" fill-opacity="{.9 if c is None else .55}"/>')
 if q:svg.append('<polyline points="'+' '.join(f'{x(vx):.3f},{y(vy):.3f}' for vx,vy in q)+'" fill="none" stroke="#aa4b2c" stroke-width="2.5"/>')
 svg.append('</g>')
 if c is None:svg.append('<text x="460" y="120" transform="rotate(-30.26 460 120)" font-size="20" fill="#687c76">1:1</text>')
 svg.append(f'<text x="316" y="414" text-anchor="middle" font-size="19">{html.escape(xlabel)}</text><text transform="translate(22 210) rotate(-90)" text-anchor="middle" font-size="19">{html.escape(ylabel)}</text>')
 if c is not None:svg.append('<g font-size="19"><circle cx="86" cy="30" r="4" fill="#657da9"/><text x="100" y="35">Natural data</text><path d="M270 30H294" stroke="#aa4b2c" stroke-width="2"/><text x="304" y="35">Model</text></g>')
 svg.append('</g></svg>')
 (site/'assets'/f'{file}.svg').write_text(''.join(svg));print(file,len(p),len(q))

# Keep approved site copy when rebuilding from the source document.
import runpy
runpy.run_path(str(site/"scripts/polish-content.py"))
